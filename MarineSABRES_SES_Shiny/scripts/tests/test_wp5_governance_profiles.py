"""Verify data/wp5_governance_profiles.json against the D5.4 sources (spec 2026-10-17 WS-9a §3.4).

Comparison keeps only letters and digits (lower-cased), so line breaks, hyphenation,
quotes and punctuation do not matter, but every word must be the source's.
"""
import json
import os
import re
import unicodedata
from functools import lru_cache
from pathlib import Path

import pytest

APP = Path(__file__).resolve().parents[2]
WP5 = Path(os.environ.get("WP5_DIR", "C:/Users/arturas.baziukas/OneDrive - ku.lt/HORIZON_EUROPE/Marine-SABRES/"
                                     "SESToolbox/MarineSABRES_SES_Shiny/WP5"))
T54 = WP5 / "Task 5.4 Governance Measures & Targets"
D54 = T54 / "Deliverable 5.4" / "final version" / "MarineSABRES_D5.4 Governance Arrangements_29.05.2026.pdf"
CHAR = {"char_tus": T54 / "Case study work" / "Tuscany archipelago" / "MarineSABRES_Characterisation_TuscanArchipelago.docx",
        "char_mac": T54 / "Case study work" / "Macaronesia" / "MarineSABRES_Characterisation_Macaronesia.docx"}
TABLE_PAGES = {"arctic_nea": range(24, 28), "macaronesia": range(29, 32), "tuscan_archipelago": range(34, 36)}  # PDF pages
# x where the "working well" and "areas of improvement" columns start; Tables 6 and 7 have a
# narrower "Dimension" column than Table 5, so the bounds differ per table
COL_X = {"arctic_nea": (180, 335), "macaronesia": (140, 345), "tuscan_archipelago": (140, 335)}
FOOTER_Y = 770
PROFILES = json.loads((APP / "data" / "wp5_governance_profiles.json").read_text(encoding="utf-8"))
needs_src = pytest.mark.skipif(not D54.exists(), reason="D5.4 sources not available (WP5/ is not in git)")


def key(s):
    return re.sub(r"[^0-9a-z]", "", unicodedata.normalize("NFKC", s).lower())


@lru_cache(maxsize=None)
def doc():
    import fitz
    return fitz.open(str(D54))


@lru_cache(maxsize=None)
def full_text():
    # words above the footer only, so printed page numbers do not split a paragraph that runs on
    return key(" ".join(w[4] for p in doc() for w in p.get_text("words") if w[1] < FOOTER_Y))


def table_body(page):
    """(top, bottom) of the table body on a page: below the repeated header row, above the last rule."""
    ys = set()
    for d in page.get_drawings():
        for it in d["items"]:
            if it[0] == "l" and abs(it[1].y - it[2].y) < 1 and abs(it[1].x - it[2].x) > 100:
                ys.add(round(it[1].y))
            elif it[0] == "re" and it[1].height < 2 and it[1].width > 100:
                ys.add(round(it[1].y0))
    ys = sorted(ys)
    return ys[1], ys[-1]


@lru_cache(maxsize=None)
def column_text(da, polarity):
    lo, hi = (COL_X[da][0], COL_X[da][1]) if polarity == "enabler" else (COL_X[da][1], 10_000)
    parts = []
    for p in TABLE_PAGES[da]:
        top, bottom = table_body(doc()[p - 1])
        ws = [w for w in doc()[p - 1].get_text("words") if lo <= w[0] < hi and top < w[1] < bottom]
        ws.sort(key=lambda w: (round(w[1] / 3), w[0]))
        parts.append(" ".join(w[4] for w in ws))
    return key(" ".join(parts))


@lru_cache(maxsize=None)
def docx_text(name):
    import docx
    d = docx.Document(str(CHAR[name]))
    cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
    return key(" ".join([p.text for p in d.paragraphs] + cells))


def items(kind):
    return [(da, x) for da, prof in PROFILES["das"].items() for x in prof[kind]]


def test_key_ignores_layout_only():
    assert key("Transpar-\nency: “NEAFC”") == key("Transparency: NEAFC")
    assert key("a b") != key("a c")


@needs_src
@pytest.mark.parametrize("da,card", items("cards"), ids=lambda v: v.get("id", v) if isinstance(v, dict) else v)
def test_card_text_is_in_its_table_column(da, card):
    assert key(card["text"]) in column_text(da, card["polarity"]), card["id"]


@needs_src
@pytest.mark.parametrize("da,rec", items("recommendations"), ids=lambda v: v.get("id", v) if isinstance(v, dict) else v)
def test_recommendation_title_and_text_are_in_d54(da, rec):
    assert key(rec["title"]) in full_text(), rec["id"]
    assert key(rec["text"]) in full_text(), rec["id"]


@needs_src
@pytest.mark.parametrize("da,actor", items("actors"), ids=lambda v: v.get("id", v) if isinstance(v, dict) else v)
def test_actor_name_is_in_its_source(da, actor):
    src = full_text() if actor["source_doc"] == "d54" else docx_text(actor["source_doc"])
    assert key(actor["name"]) in src, actor["id"]


@needs_src
def test_a_made_up_card_is_rejected():
    assert key("Greenland publishes all fisheries data in real time") not in column_text("arctic_nea", "enabler")
