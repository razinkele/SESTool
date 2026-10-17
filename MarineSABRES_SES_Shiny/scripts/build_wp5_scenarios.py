"""Build the WP5 D5.1 scenario x worldview reference data (spec 2026-10-16 WS-8).

Usage:
    micromamba run -n shiny python scripts/build_wp5_scenarios.py [--src DIR] [--out DIR]

Defaults: --src <app>/WP5/Task 5.1 Application of Scenarios, --out <app>/data.
The WP5 source folder is not in git; the generated JSON files are committed.
"""
import argparse
import json
import re
import sys
from pathlib import Path

import docx
import openpyxl
from docx.oxml.ns import qn

APP = Path(__file__).resolve().parents[1]
D51 = "Deliverable report T5.1/Marine SABRES D5.1 scenarios report_v2 final.docx"
RT = "Deliverable report T5.1/SSP-RCP regionalisation table.docx"
RB = "Scenarios vs worldviews paper/results breakdown.xlsx"

PESTLE = ["political", "economic", "social", "technological", "legal", "environmental"]
PESTLE_LABELS = ["Political", "Economic", "Social", "Technological", "Tecnological", "Legal", "Environmental"]
SSPS = ["ssp1_26", "ssp2_45", "ssp3_70", "ssp5_85"]
WVS = ["egalitarian", "individualist", "hierarchy"]
WV_NAME = {"EGA": "egalitarian", "IND": "individualist", "HIE": "hierarchy"}
SLICES = ["general", "hrmc", "sh", "lh"]
SLICE_HEADINGS = {
    "HABITAT RESTORATION AND MARINE CONSERVATION": "hrmc",
    "SEAFOOD HARVESTING": "sh",
    "TOURISM AND LEISURE": "lh",
}
# D5.1 Box 1 (fixed values, not parsed): warming cap by 2100 and EBM conduciveness
SSP_FIXED = {
    "ssp1_26": {"ssp": "SSP1", "rcp": "2.6", "warming_cap": "<=2", "ebm_conducive": "yes"},
    "ssp2_45": {"ssp": "SSP2", "rcp": "4.5", "warming_cap": "<=3", "ebm_conducive": "somewhat"},
    "ssp3_70": {"ssp": "SSP3", "rcp": "7.0", "warming_cap": "<=4", "ebm_conducive": "no"},
    "ssp5_85": {"ssp": "SSP5", "rcp": "8.5", "warming_cap": ">4", "ebm_conducive": "no"},
}
FILL_STATE = {"0070C0": "area", "808080": "europe", "D9D9D9": "none"}
# Cells where the working spreadsheet differs from the published D5.1 tables (spec 3.2)
EXPECTED_X_DIFF = {("hierarchy", "ssp3_70", "political")}
EXPECTED_TRANSFORM_DIFF = {("hierarchy", "ssp3_70", "political"), ("hierarchy", "ssp3_70", "environmental")}
TICK, CROSS = "ü", "ý"  # Wingdings 0xFC = needed, 0xFD = not needed


def clean(s):
    s = (s or "").replace("\u00a0", " ")
    s = re.sub(r"[ \t]*\n[ \t]*", "\n", s)
    s = re.sub(r"\n{2,}", "\n", s)
    return re.sub(r"[ \t]+", " ", s).strip()


def fill(cell):
    pr = cell._tc.tcPr
    shd = pr.find(qn("w:shd")) if pr is not None else None
    val = shd.get(qn("w:fill")) if shd is not None else None
    return val.upper() if val else None


def wv_id(text):
    return WV_NAME.get(clean(text)[:3].upper())


# --- narratives: regionalisation table, table 2 ----------------------------------
def split_slices(text):
    out, current, buf = {}, "general", []
    for line in text.split("\n"):
        key = SLICE_HEADINGS.get(line.strip())
        if key:
            out[current] = "\n".join(buf).strip()
            current, buf = key, []
        else:
            buf.append(line)
    out[current] = "\n".join(buf).strip()
    return out


def gap_code(text):
    if text is None or not text.strip():
        return "no_source_text"
    if text.startswith("SSP not featured in Future MARES"):
        return "ssp_not_in_futuremares"
    if text.startswith("No specific lines in AFMEC"):
        return "no_afmec_lines"
    return None


def build_scenarios(src):
    t = docx.Document(str(src / RT)).tables[2]
    scenarios, questions = [], {}
    for k, sid in enumerate(SSPS):
        col = 2 + k
        sc = {"id": sid, **SSP_FIXED[sid],
              "name": clean(t.rows[0].cells[col].text).split("\n", 1)[-1],
              "global_summary": clean(t.rows[1].cells[col].text),
              "expanded_narrative": clean(t.rows[2].cells[col].text),
              "driver_mapping": None, "cells": {}}
        for p, pid in enumerate(PESTLE):
            row = 4 + p
            questions[pid] = clean(t.rows[row].cells[1].text)
            parts = split_slices(clean(t.rows[row].cells[col].text))
            cell = {}
            for sl in SLICES:
                txt = parts.get(sl)
                code = gap_code(txt)
                cell[sl] = {"text": None if code else txt, "gap_note": code,
                            "source_ref": f"table2:r{row}:c{col}"}
            sc["cells"][pid] = cell
        scenarios.append(sc)
    return {"version": "1.0", "source": "D5.1 (2024-02-26), SSP-RCP regionalisation table",
            "built_by": "scripts/build_wp5_scenarios.py", "pestle": PESTLE,
            "pestle_questions": questions, "slices": SLICES,
            "slice_da_focus": {"hrmc": "tuscany", "sh": "arctic", "lh": "macaronesia"},
            "scenarios": scenarios}


# --- worldview matrix: D5.1 Tables 3, 4, 5 and Appendix 1 ------------------------
def find_tables(d):
    t3 = t4 = t13 = None
    t5 = []
    for t in d.tables:
        cols = len(t.columns)
        fills = {fill(c) for r in t.rows for c in r.cells} - {None}
        if cols == 4 and clean(t.rows[0].cells[0].text).startswith("Worldview"):
            t3 = t
        elif cols == 26 and "92D050" in fills:
            t4 = t
        elif cols == 25 and {"0070C0", "D9D9D9"} <= fills:
            t5.append(t)
        elif cols == 6 and any(clean(r.cells[1].text) in PESTLE_LABELS for r in t.rows):
            t13 = t
    if not (t3 and t4 and t13 and len(t5) == 2):
        sys.exit(f"D5.1 tables not found: t3={bool(t3)} t4={bool(t4)} t5={len(t5)} appendix={bool(t13)}")
    return t3, t4, t5, t13


def table5_states(t):
    out = {}
    for r in t.rows:
        wv = wv_id(r.cells[0].text)
        if not wv:
            continue
        for s, sid in enumerate(SSPS):
            for p, pid in enumerate(PESTLE):
                out[(wv, sid, pid)] = FILL_STATE[fill(r.cells[1 + s * 6 + p])]
    return out


def build_matrix(src):
    d = docx.Document(str(src / D51))
    t3, t4, t5s, t13 = find_tables(d)
    states = table5_states(t5s[0])
    if states != table5_states(t5s[1]):
        sys.exit("check (d) failed: the two Table 5 copies differ")
    rows3 = list(t3.rows)
    wv_cols = {wv_id(rows3[0].cells[c].text): c for c in (1, 2, 3)}
    worldviews = []
    for wv in WVS:
        c = wv_cols[wv]
        get = lambda i, c=c: clean(rows3[i].cells[c].text)
        worldviews.append({"id": wv, "keywords": get(1), "problem": get(2), "solution": get(3),
                           "summary": get(4),
                           "pestle": {pid: (get(6 + i) or None) for i, pid in enumerate(PESTLE)},
                           "source_ref": f"D5.1 Table 3, column {c}"})
    t4rows = list(t4.rows)[2:]
    app_rows, cur = [], None
    for r in t13.rows:
        w = wv_id(r.cells[0].text) or cur
        if clean(r.cells[1].text) in PESTLE_LABELS:
            app_rows.append((w, r))
        cur = w
    cells = []
    for wv in WVS:
        r4s = [r for r in t4rows if wv_id(r.cells[0].text) == wv]
        ras = [r for w, r in app_rows if w == wv]
        for p, pid in enumerate(PESTLE):
            r4, ra = r4s[p], ras[p]
            for s, sid in enumerate(SSPS):
                c4 = r4.cells[2 + s * 6 + p]
                match = {"92D050": True, "FF0000": False}[fill(c4)]
                rat = clean(ra.cells[2 + s].text)
                cells.append({"worldview": wv, "ssp": sid, "pestle": pid,
                              "transformation": states[(wv, sid, pid)],
                              "match": match,
                              "wv_more_conducive": None if match else clean(c4.text) == "X",
                              # Appendix 1 marks some "worldview more conducive" cells with a leading '*',
                              # but not all of them (8 are missing); Table 4's "X" is the source, so strip it
                              "rationale": rat.lstrip("*").strip(),
                              "source_ref": {"table4": f"{wv}:{pid}:c{2 + s * 6 + p}",
                                             "table5": f"{wv}:c{1 + s * 6 + p}",
                                             "appendix1": f"{wv}:{pid}:c{2 + s}"}})
    return {"version": "1.0", "source": "D5.1 Tables 3, 4, 5 and Appendix 1",
            "built_by": "scripts/build_wp5_scenarios.py", "worldviews": worldviews, "cells": cells,
            "source_notes": [{"id": i, "text_key": f"modules.wp5_scenarios.source_note.{i}"} for i in
                             ("summary_hierarchy_ssp", "hierarchy_ssp2_misprint", "conducive_row_label",
                              "ssp3_hierarchy_cells", "generic_europe")],
            "strategies": ["behavioural", "accepted_instruments", "limits_of_acceptable_change"]}


# --- consistency checks against the working spreadsheet (spec 3.3) ---------------
def check_against_spreadsheet(src, matrix):
    ws = openpyxl.load_workbook(str(src / RB), data_only=True)["Sheet1"]
    rows = {"egalitarian": (5, 6, 7), "individualist": (9, 10, 11), "hierarchy": (13, 14, 15)}
    errors, x_diff, t_diff = [], set(), set()
    for c in matrix["cells"]:
        mr, xr, tr = rows[c["worldview"]]
        col = 3 + SSPS.index(c["ssp"]) * 6 + PESTLE.index(c["pestle"])
        key = (c["worldview"], c["ssp"], c["pestle"])
        if (ws.cell(mr, col).value == "Y") != c["match"]:
            errors.append(f"(a) match differs at {key}")
        if (ws.cell(xr, col).value == "Y") != bool(c["wv_more_conducive"]):
            x_diff.add(key)
        glyph = str(ws.cell(tr, col).value or "")[:1]
        if {TICK: True, CROSS: False}.get(glyph) != (c["transformation"] != "none"):
            t_diff.add(key)
    if x_diff != EXPECTED_X_DIFF:
        errors.append(f"(b) unexpected more-conducive differences: {sorted(x_diff ^ EXPECTED_X_DIFF)}")
    if t_diff != EXPECTED_TRANSFORM_DIFF:
        errors.append(f"(c) unexpected transformation differences: {sorted(t_diff ^ EXPECTED_TRANSFORM_DIFF)}")
    return errors


def build(src, out):
    scen = build_scenarios(src)
    matrix = build_matrix(src)
    errors = check_against_spreadsheet(src, matrix)
    if errors:
        print("\n".join(errors))
        sys.exit(1)
    out.mkdir(parents=True, exist_ok=True)
    for name, obj in (("wp5_scenarios_ssp.json", scen), ("wp5_worldview_matrix.json", matrix)):
        (out / name).write_text(json.dumps(obj, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
                                encoding="utf-8", newline="\n")
    return scen, matrix


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--src", default=str(APP / "WP5" / "Task 5.1 Application of Scenarios"))
    ap.add_argument("--out", default=str(APP / "data"))
    a = ap.parse_args(argv)
    build(Path(a.src), Path(a.out))
    print(f"wrote wp5_scenarios_ssp.json and wp5_worldview_matrix.json to {a.out}")


if __name__ == "__main__":
    main()
